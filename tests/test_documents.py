import subprocess
from pathlib import Path

from docx import Document
from pypdf import PdfWriter

from knowledge_server.documents import (
    discover_documents,
    discover_git_documents,
    load_document,
)


def test_load_html_removes_scripts_and_preserves_text(tmp_path: Path) -> None:
    source = tmp_path / "onenote-page.html"
    source.write_text(
        "<html><style>.hidden{}</style><h1>SQL</h1>"
        "<p>DIM_Environment</p><script>ignore()</script></html>",
        encoding="utf-8",
    )

    document = load_document(source)

    assert "SQL" in document.content
    assert "DIM_Environment" in document.content
    assert "ignore" not in document.content


def test_load_docx_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    source = tmp_path / "notes.docx"
    word_document = Document()
    word_document.add_paragraph("Architectuurnotitie")
    table = word_document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Bron"
    table.rows[0].cells[1].text = "Target"
    word_document.save(source)

    document = load_document(source)

    assert "Architectuurnotitie" in document.content
    assert "Bron | Target" in document.content


def test_load_pdf_keeps_page_markers(tmp_path: Path) -> None:
    source = tmp_path / "notes.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with source.open("wb") as output:
        writer.write(output)

    document = load_document(source)

    assert document.content == "# Pagina 1"


def test_discover_documents_supports_extended_formats(tmp_path: Path) -> None:
    for filename in (
        "Dockerfile",
        "a.md",
        "b.sql",
        "c.html",
        "d.docx",
        "e.pdf",
        "f.py",
        "g.yaml",
    ):
        (tmp_path / filename).touch()
    (tmp_path / "ignored.png").touch()

    discovered = discover_documents(tmp_path)

    assert [path.name for path in discovered] == [
        "Dockerfile",
        "a.md",
        "b.sql",
        "c.html",
        "d.docx",
        "e.pdf",
        "f.py",
        "g.yaml",
    ]


def test_discover_git_documents_only_returns_safe_tracked_files(tmp_path: Path) -> None:
    subprocess.run(["git", "init", "-q", str(tmp_path)], check=True)
    (tmp_path / "app.py").write_text("print('safe')\n", encoding="utf-8")
    (tmp_path / ".env").write_text("TOKEN=secret\n", encoding="utf-8")
    (tmp_path / "notes.tmp").write_text("not supported\n", encoding="utf-8")
    subprocess.run(
        ["git", "-C", str(tmp_path), "add", "app.py", ".env", "notes.tmp"],
        check=True,
    )

    discovered = discover_git_documents(tmp_path)

    assert [path.name for path in discovered] == ["app.py"]
