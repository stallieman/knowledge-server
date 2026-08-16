from __future__ import annotations

import hashlib
import subprocess
from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

SUPPORTED_SUFFIXES = frozenset(
    {
        ".bat",
        ".cfg",
        ".conf",
        ".cs",
        ".css",
        ".docx",
        ".fish",
        ".go",
        ".htm",
        ".html",
        ".ini",
        ".java",
        ".js",
        ".json",
        ".jsx",
        ".md",
        ".pdf",
        ".properties",
        ".ps1",
        ".py",
        ".rs",
        ".scss",
        ".sh",
        ".sql",
        ".toml",
        ".ts",
        ".tsx",
        ".txt",
        ".xml",
        ".yaml",
        ".yml",
    }
)
SUPPORTED_FILENAMES = frozenset(
    {".env.example", "containerfile", "dockerfile", "justfile", "makefile"}
)
EXCLUDED_DIRECTORIES = frozenset(
    {
        ".git",
        ".mypy_cache",
        ".pytest_cache",
        ".ruff_cache",
        ".tox",
        ".venv",
        "__pycache__",
        "build",
        "dist",
        "node_modules",
        "target",
        "vendor",
    }
)
SENSITIVE_FILENAMES = frozenset(
    {
        ".env",
        "credentials",
        "credentials.json",
        "id_dsa",
        "id_ecdsa",
        "id_ed25519",
        "id_rsa",
        "secrets.json",
    }
)
SENSITIVE_SUFFIXES = frozenset({".key", ".p12", ".pem", ".pfx"})
GENERATED_FILENAMES = frozenset(
    {"cargo.lock", "package-lock.json", "pnpm-lock.yaml", "yarn.lock"}
)
MAX_TEXT_DOCUMENT_BYTES = 2_000_000
MAX_BINARY_DOCUMENT_BYTES = 50_000_000


@dataclass(frozen=True)
class LoadedDocument:
    """Text loaded from one local source file."""

    source_path: Path
    content: str
    content_hash: str


def discover_documents(path: Path) -> list[Path]:
    """Find supported documents at a path in deterministic order."""
    resolved_path = path.expanduser().resolve()
    if not resolved_path.exists():
        raise FileNotFoundError(f"Source path does not exist: {resolved_path}")

    if resolved_path.is_file():
        candidates = [resolved_path]
    else:
        candidates = [
            candidate for candidate in resolved_path.rglob("*") if candidate.is_file()
        ]

    return sorted(
        candidate
        for candidate in candidates
        if _is_indexable_path(candidate, root=resolved_path)
    )


def discover_git_documents(repository_path: Path) -> list[Path]:
    """List safe, supported files tracked by Git without reading ignored files."""
    repository = repository_path.expanduser().resolve()
    if not (repository / ".git").exists():
        raise ValueError(f"Geen lokale Git-repository: {repository}")

    result = subprocess.run(
        ["git", "-C", str(repository), "ls-files", "-z"],
        check=True,
        capture_output=True,
    )
    relative_paths = [
        Path(value.decode("utf-8", errors="surrogateescape"))
        for value in result.stdout.split(b"\0")
        if value
    ]
    return sorted(
        path
        for relative_path in relative_paths
        if (path := (repository / relative_path).resolve()).is_file()
        and path.is_relative_to(repository)
        and _is_indexable_path(path, root=repository)
    )


def _is_indexable_path(path: Path, *, root: Path) -> bool:
    try:
        relative = path.relative_to(root)
    except ValueError:
        return False

    if any(part in EXCLUDED_DIRECTORIES for part in relative.parts[:-1]):
        return False
    filename = path.name.casefold()
    if filename in SENSITIVE_FILENAMES or path.suffix.casefold() in SENSITIVE_SUFFIXES:
        return False
    if filename in GENERATED_FILENAMES or filename.endswith("-lock.json"):
        return False
    if not (
        path.suffix.casefold() in SUPPORTED_SUFFIXES or filename in SUPPORTED_FILENAMES
    ):
        return False
    try:
        size = path.stat().st_size
    except OSError:
        return False
    size_limit = (
        MAX_BINARY_DOCUMENT_BYTES
        if path.suffix.casefold() in {".docx", ".pdf"}
        else MAX_TEXT_DOCUMENT_BYTES
    )
    return size <= size_limit


def load_document(path: Path) -> LoadedDocument:
    """Load text from a supported local document."""
    resolved_path = path.expanduser().resolve()
    suffix = resolved_path.suffix.lower()
    if (
        suffix not in SUPPORTED_SUFFIXES
        and resolved_path.name.casefold() not in SUPPORTED_FILENAMES
    ):
        raise ValueError(f"Unsupported document type: {suffix}")

    if suffix == ".pdf":
        content = _load_pdf(resolved_path)
    elif suffix == ".docx":
        content = _load_docx(resolved_path)
    elif suffix in {".htm", ".html"}:
        content = _load_html(resolved_path)
    else:
        content = resolved_path.read_text(encoding="utf-8", errors="replace")

    content_hash = hashlib.sha256(content.encode("utf-8")).hexdigest()
    return LoadedDocument(
        source_path=resolved_path,
        content=content,
        content_hash=content_hash,
    )


def _load_pdf(path: Path) -> str:
    pages = []
    for page_number, page in enumerate(PdfReader(path).pages, start=1):
        page_text = page.extract_text() or ""
        pages.append(f"# Pagina {page_number}\n\n{page_text.strip()}")
    return "\n\n".join(pages).strip()


def _load_docx(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs]

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            parts.append(" | ".join(values))

    return "\n".join(part for part in parts if part).strip()


def _load_html(path: Path) -> str:
    soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
    for element in soup(["script", "style"]):
        element.decompose()
    return soup.get_text("\n", strip=True)
